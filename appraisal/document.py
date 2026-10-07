"""
Exports a board-ready business case as a Word document, structured on the
UK Government Five Case Model (HM Treasury Green Book):

  1. Strategic case   - why change is needed
  2. Economic case    - which option gives best value (NPV, BCR, risk)
  3. Commercial case  - how it will be procured
  4. Financial case   - is it affordable, and when is the cash needed
  5. Management case  - how delivery and benefits will be governed
"""

from __future__ import annotations

import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from .model import Assumptions, Option, appraise, cash_flows, scenarios, tornado, value_for_money
from .report import money, years
from .simulation import Uncertainty, cdf, simulate, summarise

NAVY = RGBColor(0x10, 0x42, 0x81)
BLUE, ORANGE, AQUA, GREY = "#2a78d6", "#eb6834", "#1baf7a", "#9a9993"

DEFAULT_NARRATIVE = {
    "strategic": "Monthly status reporting across the 40-project portfolio is compiled by hand from six "
                 "different trackers. It absorbs roughly 80 analyst-days a month, reports are a week old by the "
                 "time the board reads them, and figures regularly disagree between packs. Leadership wants "
                 "decisions made on current, single-source data.",
    "objectives": "Cut reporting effort by at least 60%; deliver reports within 2 working days of month end; "
                  "one source of truth for cost, schedule, risk and resourcing data.",
    "commercial": "Procure via an existing public-sector framework (e.g. G-Cloud) to shorten the timeline. "
                  "Three-year term with a two-year extension option; data export rights and exit support "
                  "written into the contract.",
    "management": "Senior Responsible Owner: Programme Director. Delivery via a 6-month phased rollout with a "
                  "pilot on 5 projects. Benefits owned by named individuals and reviewed quarterly at the "
                  "portfolio board; the business case is re-validated at each stage gate.",
}


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "104281")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph()


def _figure(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _axes(ax) -> None:
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#bdbcb6"); ax.spines["bottom"].set_color("#bdbcb6")
    ax.tick_params(colors="#52514e", labelsize=8)
    ax.grid(axis="y", color="#e6e5e0", linewidth=0.8)
    ax.set_axisbelow(True)


def _gbp_axis(ax, axis: str = "y") -> None:
    fmt = matplotlib.ticker.FuncFormatter(lambda v, _: money(v))
    (ax.yaxis if axis == "y" else ax.xaxis).set_major_formatter(fmt)


def build(title: str, options: list[Option], a: Assumptions, chosen: str,
          narrative: dict | None = None, unc: Uncertainty = Uncertainty(), author: str = "") -> bytes:
    n = {**DEFAULT_NARRATIVE, **(narrative or {})}
    res = pd.DataFrame([appraise(o, a) for o in options])
    opt = next(o for o in options if o.name == chosen)
    pick = res.set_index("option").loc[chosen]
    sims = simulate(options, a, unc)
    odds = summarise(sims).set_index("option")
    sc = scenarios(opt, a).set_index("scenario")
    tor = tornado(opt, a).sort_values("range", ascending=False)
    cf = cash_flows(opt, a)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    for lvl in (1, 2):
        doc.styles[f"Heading {lvl}"].font.color.rgb = NAVY

    h = doc.add_heading(f"Business case: {title}", 0)
    h.runs[0].font.color.rgb = NAVY
    meta = doc.add_paragraph(f"Structured on the HM Treasury Five Case Model. Recommended option: {chosen}.")
    if author:
        meta.add_run(f"  Prepared by {author}.")
    meta.runs[0].italic = True

    doc.add_heading("Executive summary", 1)
    p = doc.add_paragraph()
    p.add_run("Recommendation: ").bold = True
    p.add_run(f"approve {chosen}. Over {a.years} years it delivers a net present value of {money(pick.npv)} "
              f"and a benefit-cost ratio of {pick.bcr:.2f} ({value_for_money(pick.bcr).lower()} value for money), "
              f"paying back in {years(pick.payback_years)}.")
    p = doc.add_paragraph()
    p.add_run("Risk: ").bold = True
    p.add_run(f"across 10,000 simulated outcomes it pays for itself {odds.loc[chosen, 'p_positive']:.0%} of the time "
              f"and is the best option in {odds.loc[chosen, 'p_best']:.0%} of them. The middle outcome (P50) is "
              f"{money(odds.loc[chosen, 'p50'])}; a bad outcome (P10) is {money(odds.loc[chosen, 'p10'])}.")
    p = doc.add_paragraph()
    p.add_run("Key sensitivity: ").bold = True
    p.add_run(f"the answer depends most on {tor.iloc[0].driver.lower()}. Benefits could fall "
              f"{pick.benefit_headroom:.0%} before the option stops paying for itself.")

    doc.add_heading("1. Strategic case", 1)
    doc.add_heading("Case for change", 2)
    doc.add_paragraph(n["strategic"])
    doc.add_heading("Objectives", 2)
    doc.add_paragraph(n["objectives"])

    doc.add_heading("2. Economic case", 1)
    doc.add_paragraph(f"Options are appraised over {a.years} years at a {a.discount_rate:.1%} discount rate. "
                      f"Costs carry a {a.cost_optimism_bias:.0%} optimism-bias uplift (plus option-specific "
                      f"delivery-risk uplifts) and benefits a {a.benefit_haircut:.0%} haircut, in line with Green "
                      "Book guidance that early estimates are systematically optimistic.")
    rows = []
    for r in res.itertuples():
        base = r.pv_costs == 0
        rows.append([r.option, money(r.pv_costs), money(r.pv_benefits), money(r.npv),
                     "baseline" if base else f"{r.bcr:.2f}", "-" if base else value_for_money(r.bcr),
                     "-" if base else years(r.payback_years)])
    _table(doc, ["Option", "PV costs", "PV benefits", "NPV", "BCR", "Value for money", "Payback"], rows)

    fig, ax = plt.subplots(figsize=(6.4, 2.6))
    ax.plot(cf.year, cf.cumulative_net, color=BLUE, marker="o", lw=2, label="Cumulative (cash)")
    ax.plot(cf.year, cf.cumulative_pv_net, color=BLUE, marker="o", lw=2, ls=":", label="Cumulative (discounted)")
    ax.axhline(0, color="#52514e", lw=0.8)
    _axes(ax); _gbp_axis(ax)
    ax.set_xlabel("Year", fontsize=8, color="#52514e")
    ax.set_title(f"{chosen}: when it pays for itself", fontsize=10, loc="left")
    ax.legend(frameon=False, fontsize=8)
    doc.add_picture(_figure(fig), width=Cm(15.5))

    doc.add_heading("Risk analysis", 2)
    rows = [[o, money(r.p10), money(r.p50), money(r.p90), f"{r.p_positive:.0%}",
             "-" if pd.isna(r.p_best) else f"{r.p_best:.0%}"]
            for o, r in odds.iterrows() if res.set_index("option").loc[o, "pv_costs"] > 0]
    _table(doc, ["Option", "P10 NPV", "P50 NPV", "P90 NPV", "Chance NPV > 0", "Chance best option"], rows)
    fig, ax = plt.subplots(figsize=(6.4, 2.6))
    palette = [BLUE, ORANGE, AQUA, "#4a3aa7"]
    for i, o in enumerate([o for o in sims.columns if res.set_index("option").loc[o, "pv_costs"] > 0]):
        c = cdf(sims[o])
        ax.plot(c.npv, c.p_at_least, lw=2, color=palette[i % len(palette)], label=o)
    ax.axvline(0, color="#52514e", lw=0.8)
    _axes(ax); _gbp_axis(ax, "x")
    ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    ax.set_title("Probability NPV is at least X", fontsize=10, loc="left")
    ax.legend(frameon=False, fontsize=8)
    doc.add_picture(_figure(fig), width=Cm(15.5))
    doc.add_paragraph(f"Simulation ranges: costs {unc.cost_low:+.0%} to {unc.cost_high:+.0%}, benefits "
                      f"{unc.benefit_low:+.0%} to {unc.benefit_high:+.0%}, {unc.delay_probability:.0%} chance benefits "
                      "start a year late.").runs[0].italic = True

    doc.add_heading("Sensitivity and scenarios", 2)
    _table(doc, ["Driver (±20%)", "NPV if lower", "NPV if higher", "Swing"],
           [[r.driver, money(r.npv_low_input), money(r.npv_high_input), money(r.range)] for r in tor.itertuples()])
    _table(doc, ["Scenario", "NPV", "BCR", "Payback"],
           [[s, money(r.npv), f"{r.bcr:.2f}", years(r.payback_years)] for s, r in sc.iterrows()])

    doc.add_heading("3. Commercial case", 1)
    doc.add_paragraph(n["commercial"])

    doc.add_heading("4. Financial case", 1)
    doc.add_paragraph("Cash required by year for the recommended option (including optimism bias):")
    _table(doc, ["Year", "Costs", "Benefits", "Net", "Cumulative"],
           [[int(r.year), money(r.costs), money(r.benefits), money(r.net), money(r.cumulative_net)]
            for r in cf.itertuples()])
    peak = cf.cumulative_net.min()
    doc.add_paragraph(f"Peak funding requirement is {money(-peak)} (the lowest point of cumulative cash flow), "
                      "which is the budget that must be secured up front.")

    doc.add_heading("5. Management case", 1)
    doc.add_paragraph(n["management"])
    doc.add_paragraph("Benefits will be tracked quarterly against this case. If the re-forecast NPV turns "
                      "negative, the project will be returned to the board for a stop / continue decision "
                      "(continued business justification).")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
